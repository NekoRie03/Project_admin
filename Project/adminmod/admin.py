from django.contrib.admin.models import LogEntry, CHANGE
from django.contrib.contenttypes.models import ContentType
from django.utils import timezone
from django.contrib import admin
from django.utils.html import format_html
from django.urls import reverse, path
from django.db import models
from .models import User, StudentRegistration, Program, Section, Violation, Sanction, ViolationRecord
from .forms import StudentRegistrationAdminForm, StaffSignupForm
from django.contrib.auth.models import Group
from django.forms import TextInput, EmailInput, PasswordInput
from import_export.admin import ImportExportModelAdmin
from import_export.admin import ExportMixin
from unfold.contrib.import_export.forms import ExportForm, ImportForm, SelectableFieldsExportForm
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.contrib.admin import register
from unfold.admin import ModelAdmin
from unfold.forms import AdminPasswordChangeForm, UserChangeForm, UserCreationForm
from unfold.admin import StackedInline, TabularInline
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from django.http import HttpResponse, JsonResponse
from django.db.models import Q
from datetime import timedelta, date
from django.utils.timezone import now
from django.contrib import messages
from django.shortcuts import get_object_or_404, render, redirect
from django.template.response import TemplateResponse
from django.middleware import csrf
from django.utils.html import format_html
from django.middleware.csrf import get_token





# Set Admin Header
admin.site.site_header = "Student Violation System Administration"
admin.site.site_title = "Student Violation System Admin Portal"
admin.site.index_title = "Welcome to Student Violation System Portal"
admin.site.unregister(Group)

class LogUtils:
    @staticmethod
    def create_log_entry(request_user, obj, action):
        """
        Create a log entry for administrative actions
        
        :param request_user: The user performing the action
        :param obj: The object being modified
        :param action: Description of the action taken
        """
        LogEntry.objects.log_action(
            user_id=request_user.id,
            content_type_id=ContentType.objects.get_for_model(obj).id,
            object_id=obj.id,
            object_repr=str(obj),
            action_flag=CHANGE,
            change_message=action
        )

class ApprovalStatusFilter(admin.SimpleListFilter):
    title = 'Approval Status'
    parameter_name = 'approval_status'

    def lookups(self, request, model_admin):
        return (
            (None, 'Pending'),
            ('approved', 'Approved'),
            ('rejected', 'Rejected')
        )

    def queryset(self, request, queryset):
        if self.value() == 'approved':
            return queryset.filter(is_approved=True)
        if self.value() == 'rejected':
            return queryset.filter(is_approved=False)
        if self.value() is None:
            return queryset.filter(is_approved=None)

@admin.register(StudentRegistration)
class StudentRegistrationAdmin(ExportMixin, ModelAdmin):
    list_per_page = 50
    list_max_show_all = 500
    form = StudentRegistrationAdminForm

    list_display = (
        'display_username',
        'full_name',
        'view_program',
        'view_section',
        'view_registration_date',
        'view_cor',
        'view_id',
        'actions_buttons',
        'approval_status',
        'assign_qr_code_button'
    )
    list_filter = (
        'is_approved',
        'registration_date',
        'review_date',
        'program',
        'section'
    )
    search_fields = (
        'user__username',
        'user__first_name',
        'user__last_name',
        'program__name',
        'program__code',
        'section__name'
    )
    readonly_fields = ('registration_date', 'review_date')
    actions = ['approve_selected', 'reject_selected', 'export_as_pdf']

    def display_username(self, obj):
        return obj.user.username if obj.user else "No User"
    display_username.short_description = 'Username'
    display_username.admin_order_field = 'user__username'

    def full_name(self, obj):
        return f"{obj.user.first_name} {obj.user.last_name}" if obj.user else "No User"
    full_name.short_description = 'Full Name'
    full_name.admin_order_field = 'user__first_name'

    def view_program(self, obj):
        return obj.program.name if obj.program else "Not Assigned"
    view_program.short_description = 'Program'
    view_program.admin_order_field = 'program__name'

    def view_section(self, obj):
        return obj.section.name if obj.section else "Not Assigned"
    view_section.short_description = 'Section'
    view_section.admin_order_field = 'section__name'

    def view_registration_date(self, obj):
        return obj.registration_date.strftime('%Y-%m-%d %H:%M:%S') if obj.registration_date else "No date"
    view_registration_date.short_description = 'Registration Date'
    view_registration_date.admin_order_field = 'registration_date'

    def view_cor(self, obj):
        return self._render_image(obj, 'cor_image')
    view_cor.short_description = 'COR'

    def view_id(self, obj):
        return self._render_image(obj, 'id_image')
    view_id.short_description = 'ID'

    def _render_image(self, obj, image_field_name):
        image_field = getattr(obj, image_field_name, None)
        if image_field:
            return format_html(
                '<a href="{}" target="_blank">'
                '<img src="{}" style="max-height: 50px;"/></a>',
                image_field.url, image_field.url
            )
        return "No image"

    def actions_buttons(self, obj):
        if obj.is_approved is None:
            approve_url = reverse('admin:adminmod_studentregistration_change', args=[obj.pk]) + '?action=approve'
            reject_url = reverse('admin:adminmod_studentregistration_change', args=[obj.pk]) + '?action=reject'
            return format_html(
                '<a class="button" href="{}" style="background-color: #4CAF50; color: white; margin-right: 5px;">Approve</a>'
                '<a class="button" href="{}" style="background-color: #f44336; color: white;">Reject</a>',
                approve_url, reject_url
            )
        return "Processed"
    actions_buttons.short_description = 'Actions'

    def approval_status(self, obj):
        status_colors = {
            None: '#FFA500',  # Pending
            True: '#008000',  # Approved
            False: '#FF0000'  # Rejected
        }
        status_text = {None: 'Pending', True: 'Approved', False: 'Rejected'}
        return format_html(
            '<span style="color: {};">{}</span>',
            status_colors[obj.is_approved],
            status_text[obj.is_approved]
        )
    approval_status.short_description = 'Status'

    def export_as_pdf(self, request, queryset):
        """Export selected registrations as a PDF file"""
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="student_registrations.pdf"'

        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=letter)

        pdf.drawString(100, 750, "Student Registrations Report")
        y = 730

        for obj in queryset:
            pdf.drawString(100, y, f"Username: {obj.user.username if obj.user else 'No User'}")
            pdf.drawString(100, y - 20, f"Full Name: {obj.user.get_full_name() if obj.user else 'No User'}")
            pdf.drawString(100, y - 40, f"Program: {obj.program.name if obj.program else 'Not Assigned'}")
            pdf.drawString(100, y - 60, f"Section: {obj.section.name if obj.section else 'Not Assigned'}")
            pdf.drawString(100, y - 80, f"Status: {obj.get_is_approved_display()}")
            y -= 100

            if y < 100:
                pdf.showPage()
                y = 750

        pdf.save()
        buffer.seek(0)
        response.write(buffer.read())
        return response

    def get_queryset(self, request):
        return super().get_queryset(request).select_related('user', 'program', 'section')

    def get_fieldsets(self, request, obj=None):
        if obj:  # Change view
            return (
                ('User Information', {
                    'fields': ('user', 'registration_date')
                }),
                ('Program and Section', {
                    'fields': ('program', 'section')
                }),
                ('Documents', {
                    'fields': ('cor_image', 'id_image')
                }),
                ('Review Information', {
                    'fields': ('is_approved', 'review_comments', 'review_date')
                }),
                ('Change Confirmation', {
                    'fields': ('admin_password',),
                }),
            )
        else:  # Add view
            return (
                ('User Information', {
                    'fields': ('user',)
                }),
                ('Program and Section', {
                    'fields': ('program', 'section')
                }),
                ('Documents', {
                    'fields': ('cor_image', 'id_image')
                }),
                ('Review Information', {
                    'fields': ('is_approved', 'review_comments')
                }),
                ('Change Confirmation', {
                    'fields': ('admin_password',),
                }),
            )

    def get_readonly_fields(self, request, obj=None):
        if obj and obj.is_approved is None:
            return ['registration_date', 'review_date']
        elif obj and obj.is_approved is not None:
            return ['user', 'registration_date', 'review_date', 'cor_image', 'id_image']
        return []

    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "user":
            kwargs["queryset"] = User.objects.filter(role=User.Role.STUDENT)
        elif db_field.name == "section":
            if 'program' in request.GET:
                kwargs["queryset"] = Section.objects.filter(program_id=request.GET['program'])
        return super().formfield_for_foreignkey(db_field, request, **kwargs)

    def save_model(self, request, obj, form, change):
        if change:
            obj.review_date = timezone.now()
        super().save_model(request, obj, form, change)
        LogUtils.create_log_entry(request.user, obj, "Changed via admin interface")

    def change_view(self, request, object_id, form_url='', extra_context=None):
        obj = self.get_object(request, object_id)
        if 'action' in request.GET and obj:
            action = request.GET['action']
            if action == 'approve':
                obj.approve_registration('Approved by admin')
                LogUtils.create_log_entry(request.user, obj, "Registration approved")
                self.message_user(request, 'Registration has been approved.')
            elif action == 'reject':
                obj.decline_registration('Rejected by admin')
                LogUtils.create_log_entry(request.user, obj, "Registration rejected")
                self.message_user(request, 'Registration has been rejected.')
        return super().change_view(request, object_id, form_url, extra_context)

    def _bulk_update_status(self, request, queryset, status, action_text):
        try:
            queryset.update(is_approved=status, review_date=timezone.now())
            for obj in queryset:
                LogUtils.create_log_entry(request.user, obj, f'Registration {action_text} in bulk')
            self.message_user(request, f"{queryset.count()} registrations have been {action_text}.")
        except Exception as e:
            self.message_user(request, f"An error occurred: {str(e)}", level='error')

    def approve_selected(self, request, queryset):
        self._bulk_update_status(request, queryset, True, 'approved')
    approve_selected.short_description = "Approve selected registrations"

    def reject_selected(self, request, queryset):
        self._bulk_update_status(request, queryset, False, 'rejected')
    reject_selected.short_description = "Reject selected registrations"

    def assign_qr_code_button(self, obj):
        if obj.qr_code:
            return format_html(
                '<span style="color: green;">QR Code: {}</span>',
                obj.qr_code
            )
        return format_html(
            '<a class="button" onclick="window.open(\'assign_qr_code/{}/\', \'Assign QR Code\', \'width=400,height=200\')" '
            'href="javascript:void(0)" style="background-color: #2196F3; color: white;">Assign QR Code</a>',
            obj.id
        )
    assign_qr_code_button.short_description = 'QR Code'
    assign_qr_code_button.allow_tags = True

    def assign_qr_code(self, request, student_id):
        student = get_object_or_404(StudentRegistration, pk=student_id)
        
        if request.method == "POST":
            qr_code = request.POST.get('qr_code')
            if qr_code:
                # Check if QR code is already assigned to another student
                if StudentRegistration.objects.filter(qr_code=qr_code).exclude(pk=student_id).exists():
                    return HttpResponse(
                        f'''
                        <form method="POST">
                            <input type="hidden" name="csrfmiddlewaretoken" value="{get_token(request)}">
                            <p style="color: red;">This QR Code is already assigned to another student. Please use a different code.</p>
                            <input type="text" name="qr_code" required>
                            <input type="submit" value="Assign">
                        </form>
                        '''
                    )
                
                # Assign the QR code to the student registration
                student.qr_code = qr_code
                student.save()
                
                # Update ViolationRecords that are linked to this student with the new QR code
                violation_records = ViolationRecord.objects.filter(student=student.user, qr_code__isnull=True)
                for record in violation_records:
                    record.qr_code = qr_code
                    record.save()
                
                # Log the action
                LogUtils.create_log_entry(request.user, student, f"Assigned QR code: {qr_code}")
                self.message_user(request, "QR Code assigned successfully!")
                
                # Refresh the parent window (admin interface)
                return HttpResponse(
                    '<script>window.opener.location.reload(); window.close();</script>'
                )
            else:
                return HttpResponse(
                    f'''
                    <form method="POST">
                        <input type="hidden" name="csrfmiddlewaretoken" value="{get_token(request)}">
                        <p style="color: red;">QR Code is required. Please enter a QR Code.</p>
                        <input type="text" name="qr_code" required>
                        <input type="submit" value="Assign">
                    </form>
                    '''
                )
        
        return HttpResponse(
            f'''
            <form method="POST">
                <input type="hidden" name="csrfmiddlewaretoken" value="{get_token(request)}">
                <p>Please enter QR code to assign:</p>
                <input type="text" name="qr_code" required autofocus>
                <input type="submit" value="Assign">
            </form>
            '''
        )

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'assign_qr_code/<int:student_id>/',
                self.admin_site.admin_view(self.assign_qr_code),
                name='student-registration-assign-qr-code',
            ),
        ]
        return custom_urls + urls

class SectionInline(TabularInline):
    model = Section
    extra = 1
    show_change_link = True



@admin.register(Program)
class ProgramAdmin(ImportExportModelAdmin, ModelAdmin):
    list_display = ('name', 'code', 'section_count')
    search_fields = ('name', 'code')
    list_filter = ('name',)
    inlines = [SectionInline]
    
    def section_count(self, obj):
        return obj.sections.count()
    section_count.short_description = 'Number of Sections'
    import_form_class = ImportForm
    export_form_class = ExportForm

@admin.register(Section)
class SectionAdmin(ImportExportModelAdmin, ModelAdmin):
    list_display = ('name', 'program', 'program_code')
    search_fields = ('name', 'program__name', 'program__code')
    list_filter = ('program',)
    
    def program_code(self, obj):
        return obj.program.code
    program_code.short_description = 'Program Code'
    
    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "program":
            kwargs["queryset"] = Program.objects.all().order_by('name')
        return super().formfield_for_foreignkey(db_field, request, **kwargs)
    import_form_class = ImportForm
    export_form_class = ExportForm

class SanctionInline(StackedInline):
    model = Sanction
    extra = 1
    show_change_link = True

@admin.register(Violation)
class ViolationAdmin(ImportExportModelAdmin, ModelAdmin):
    list_display = (
        'name', 
        'severity_colored', 
        'sanction_count', 
        'brief_description'
    )
    search_fields = ('name', 'description')
    list_filter = ('severity',)
    inlines = [SanctionInline]
    
    def severity_colored(self, obj):
        severity_colors = {
            'low': 'green',
            'medium': 'orange',
            'high': 'red',
            'critical': 'darkred'
        }
        color = severity_colors.get(obj.severity, 'black')
        return format_html(
            '<span style="color: {}; font-weight: bold;">{}</span>',
            color,
            obj.get_severity_display()
        )
    severity_colored.short_description = 'Severity'
    
    def sanction_count(self, obj):
        return obj.sanctions.count()
    sanction_count.short_description = 'Sanctions'
    
    def brief_description(self, obj):
        return obj.description[:50] + '...' if obj.description and len(obj.description) > 50 else obj.description or 'No description'
    brief_description.short_description = 'Description'
    import_form_class = ImportForm
    export_form_class = ExportForm

@admin.register(Sanction)
class SanctionAdmin(ImportExportModelAdmin, ModelAdmin):
    list_display = (
        'name', 
        'violation_display', 
        'duration_display', 
        'violation_severity_colored'
    )
    search_fields = ('name', 'violation__name')
    list_filter = ('violation', 'duration_unit')
    
    def violation_display(self, obj):
        return obj.violation.name
    violation_display.short_description = 'Related Violation'
    
    def duration_display(self, obj):
        return f"{obj.duration_value} {obj.get_duration_unit_display()}"
    duration_display.short_description = 'Duration'
    
    def violation_severity_colored(self, obj):
        severity_colors = {
            'low': 'green',
            'medium': 'orange',
            'high': 'red',
            'critical': 'darkred'
        }
        severity = obj.violation.severity
        color = severity_colors.get(severity, 'black')
        return format_html(
            '<span style="color: {}; font-weight: bold;">{}</span>',
            color,
            obj.violation.get_severity_display()
        )
    violation_severity_colored.short_description = 'Violation Severity'
    
    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "violation":
            kwargs["queryset"] = Violation.objects.all().order_by('-severity', 'name')
        return super().formfield_for_foreignkey(db_field, request, **kwargs)
    import_form_class = ImportForm
    export_form_class = ExportForm
    
@admin.register(User)
class UserAdmin(BaseUserAdmin, ModelAdmin):
    form = UserChangeForm
    add_form = UserCreationForm
    change_password_form = AdminPasswordChangeForm

    formfield_overrides = {
        'username': {'widget': TextInput(attrs={'class': 'vTextField', 'placeholder': 'Enter Username'})},
        'email': {'widget': EmailInput(attrs={'class': 'vTextField', 'placeholder': 'Enter Email'})},
        'password': {'widget': PasswordInput(attrs={'class': 'vTextField'})},
    }

    fieldsets = (
        (None, {
            'fields': ('username', 'password'),
            'classes': ('card',),
        }),
        ('Personal Info', {
            'fields': ('first_name', 'last_name', 'email'),
            'classes': ('card',),
        }),
        ('Important dates', {
            'fields': ('last_login', 'date_joined'),
            'classes': ('card',),
        }),
        ('Additional Information', {
            'fields': ('role',),
            'classes': ('card',),
        }),
    )

    list_display = ('username', 'email', 'first_name', 'last_name', 'role', 'is_staff')
    
    def get_queryset(self, request):
        return super().get_queryset(request).exclude(role=User.Role.STUDENT)
    
    list_filter = ('role', 'is_staff')
    search_fields = ('username', 'email', 'first_name', 'last_name')

    add_form = StaffSignupForm
    add_fieldsets = (
        (None, {
            'classes': ('wide', 'card'),
            'fields': (
                'username', 'first_name', 'last_name',
                'email', 'employee_id', 'role',
                'password1', 'password2',
            ),
        }),
    )

    def get_form(self, request, obj=None, **kwargs):
        """
        Use special form during user creation
        """
        defaults = {}
        if obj is None:
            defaults['form'] = self.add_form
        defaults.update(kwargs)
        return super().get_form(request, obj, **defaults)

    def get_fieldsets(self, request, obj=None):
        if not obj:
            return self.add_fieldsets
        return super().get_fieldsets(request, obj)


    def get_fields(self, request, obj=None):
        if not obj:
            return list(self.add_fieldsets[0][1]['fields'])
        return super().get_fields(request, obj)
    
@admin.register(ViolationRecord)
class ViolationRecordAdmin(ModelAdmin):
    list_display = ('student', 'violation', 'sanction', 'recorded_by', 'recorded_at', 'total_hours_complied', 'status', 'view_qr_code', 'time_in_button', 'time_out_button',)
    search_fields = (
        'student__username',
        'student__first_name',
        'student__last_name',
        'violation__name',
        'sanction__name',
    )
    list_filter = ('recorded_at', 'violation__severity', 'sanction')
    readonly_fields = ('student', 'recorded_by')

    actions = [
        'export_weekly_violations_as_pdf', 
        'export_monthly_violations_as_pdf',
        'export_weekly_violations_as_excel',
        'export_monthly_violations_as_excel',
        'export_weekly_violations_as_word',
        'export_monthly_violations_as_word'
    ]

    def export_weekly_violations_as_pdf(self, request, queryset):
        """Export a PDF report of violations committed in the last week."""
        one_week_ago = now() - timedelta(days=7)
        violations_this_week = ViolationRecord.objects.filter(recorded_at__gte=one_week_ago)

        if not violations_this_week.exists():
            messages.warning(request, "No violations found for the last week.")
            return

        response = self.generate_pdf_report(
            violations=violations_this_week,
            title="Weekly Student Violations Report",
            filename="weekly_violations_report.pdf",
            date_range=(one_week_ago, now()),
        )
        # Add a success message
        messages.success(request, "Weekly violations report has been exported successfully as a PDF.")
        return response


    def export_monthly_violations_as_pdf(self, request, queryset):
        """Export a PDF report of violations committed in the last month."""
        one_month_ago = now() - timedelta(days=30)
        violations_this_month = ViolationRecord.objects.filter(recorded_at__gte=one_month_ago)

        if not violations_this_month.exists():
            messages.warning(request, "No violations found for the last month.")
            return

        response = self.generate_pdf_report(
            violations=violations_this_month,
            title="Monthly Student Violations Report",
            filename="monthly_violations_report.pdf",
            date_range=(one_month_ago, now()),
        )
        # Add a success message
        messages.success(request, "Monthly violations report has been exported successfully as a PDF.")
        return response

    def generate_pdf_report(self, violations, title, filename, date_range):
        """Generate a row-based PDF report for the given violations."""
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=letter)

        # Header section
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(175, 750, title)
        pdf.setFont("Helvetica", 12)
        pdf.drawString(175, 730, f"Report generated on: {timezone.now().strftime('%Y-%m-%d')}")

        y = 710
        pdf.drawString(
            175, y, 
            f"Violations from: {date_range[0].strftime('%Y-%m-%d')} to {date_range[1].strftime('%Y-%m-%d')}"
        )

        # Add summary statistics
        total_violations = violations.count()
        severity_counts = violations.values('violation__severity').annotate(count=models.Count('id'))

        y -= 40
        pdf.drawString(50, y, f"Total Violations: {total_violations}")

        y -= 20
        pdf.drawString(50, y, "Violations by Severity:")
        for severity_count in severity_counts:
            y -= 20
            severity = dict(Violation.SEVERITY_CHOICES).get(severity_count['violation__severity'], "Unknown")
            pdf.drawString(50, y, f"{severity}: {severity_count['count']}")

        y -= 40  # Spacing for the detailed report content

        # Loop through the violation records and display each as a row
        for record in violations:
            if y < 100:  # Start a new page if space runs out
                pdf.showPage()
                pdf.setFont("Helvetica", 11)
                y = 750

            student_name = f"{record.student.first_name} {record.student.last_name}"
            violation_name = record.violation.name
            date_recorded = record.recorded_at.strftime('%Y-%m-%d')
            severity = record.violation.get_severity_display()
            recorded_by = record.recorded_by.get_full_name() if record.recorded_by else "Unknown"
            recorded_at = record.recorded_at.strftime('%Y-%m-%d %H:%M:%S')

            # Display each field as a row
            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Student:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, student_name)
            y -= 20

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Violation:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, violation_name)
            y -= 20

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Date:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, date_recorded)
            y -= 20

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Severity:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, severity)
            y -= 20

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Recorded By:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, recorded_by)
            y -= 20

            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(50, y, "Recorded At:")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(150, y, recorded_at)
            y -= 40  # Add extra spacing between records

        pdf.save()
        buffer.seek(0)
        response.write(buffer.read())
        return response

    def export_weekly_violations_as_excel(self, request, queryset):
        """Export weekly violations to an Excel file."""
        one_week_ago = now() - timedelta(days=7)
        violations_this_week = ViolationRecord.objects.filter(recorded_at__gte=one_week_ago)

        if not violations_this_week.exists():
            messages.warning(request, "No violations found for the last week.")
            return

        return self.generate_excel_report(
            request,
            violations=violations_this_week,
            filename="weekly_violations_report.xlsx",
            title="Weekly Violations Report"
        )

    def export_monthly_violations_as_excel(self, request, queryset):
        """Export monthly violations to an Excel file."""
        one_month_ago = now() - timedelta(days=30)
        violations_this_month = ViolationRecord.objects.filter(recorded_at__gte=one_month_ago)

        if not violations_this_month.exists():
            messages.warning(request, "No violations found for the last month.")
            return

        return self.generate_excel_report(
            request,
            violations=violations_this_month,
            filename="monthly_violations_report.xlsx",
            title="Monthly Violations Report"
        )

    def export_weekly_violations_as_word(self, request, queryset):
        """Export weekly violations to a Word document."""
        one_week_ago = now() - timedelta(days=7)
        violations_this_week = ViolationRecord.objects.filter(recorded_at__gte=one_week_ago)

        if not violations_this_week.exists():
            messages.warning(request, "No violations found for the last week.")
            return

        return self.generate_word_report(
            request,
            violations=violations_this_week,
            filename="weekly_violations_report.docx",
            title="Weekly Violations Report"
        )

    def export_monthly_violations_as_word(self, request, queryset):
        """Export monthly violations to a Word document."""
        one_month_ago = now() - timedelta(days=30)
        violations_this_month = ViolationRecord.objects.filter(recorded_at__gte=one_month_ago)

        if not violations_this_month.exists():
            messages.warning(request, "No violations found for the last month.")
            return

        return self.generate_word_report(
            request,
            violations=violations_this_month,
            filename="monthly_violations_report.docx",
            title="Monthly Violations Report"
        )

    def generate_excel_report(self, request, violations, filename, title):
        """Generate an Excel report for the given violations."""
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from django.http import HttpResponse

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = title

        # Title and Date Range
        start_date = min(violations.values_list('recorded_at', flat=True))
        end_date = max(violations.values_list('recorded_at', flat=True))
    
        sheet['A1'] = title
        sheet['A1'].font = Font(bold=True, size=16)
        sheet['A2'] = f"Report generated on: {timezone.now().strftime('%Y-%m-%d')}"
        sheet['A3'] = f"Violations from: {start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}"

        # Summary Statistics
        sheet['A5'] = f"Total Violations: {violations.count()}"
        sheet['A5'].font = Font(bold=True)

        current_row = 6
        sheet['A6'] = "Violations by Severity:"
        sheet['A6'].font = Font(bold=True)

        severity_counts = violations.values('violation__severity').annotate(count=models.Count('id'))
        for severity_count in severity_counts:
            current_row += 1
            severity = dict(Violation.SEVERITY_CHOICES).get(severity_count['violation__severity'], "Unknown")
            sheet[f'A{current_row}'] = f"{severity}: {severity_count['count']}"

        # Add spacing
        current_row += 2

        # Detailed Records Header
        headers = ['Student', 'Violation', 'Date', 'Severity', 'Recorded By', 'Recorded At']
        header_row = current_row
        for col_num, header in enumerate(headers, 1):
            cell = sheet.cell(row=header_row, column=col_num)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # Detailed Records
        for record in violations:
            current_row += 1
            student_name = f"{record.student.first_name} {record.student.last_name}" if record.student else "Unknown"
            violation_name = record.violation.name if record.violation else "No Violation"
            date_recorded = record.recorded_at.strftime('%Y-%m-%d')
            severity = record.violation.get_severity_display() if record.violation else "Unknown"
            recorded_by = record.recorded_by.get_full_name() if record.recorded_by else "Unknown"
            recorded_at = record.recorded_at.strftime('%Y-%m-%d %H:%M:%S')

            row_data = [student_name, violation_name, date_recorded, severity, recorded_by, recorded_at]
            for col_num, value in enumerate(row_data, 1):
                cell = sheet.cell(row=current_row, column=col_num)
                cell.value = value

        # Adjust column widths
        for column in sheet.columns:
            max_length = 0
            column = list(column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            sheet.column_dimensions[column[0].column_letter].width = adjusted_width

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        workbook.save(response)
        messages.success(request, f"{title} has been exported as an Excel file.")
        return response

    def generate_word_report(self, request, violations, filename, title):
        """Generate a Word report for the given violations."""
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from django.http import HttpResponse

        document = Document()

        # Title
        title_paragraph = document.add_heading(title, level=1)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Date Range
        start_date = min(violations.values_list('recorded_at', flat=True))
        end_date = max(violations.values_list('recorded_at', flat=True))
    
        date_para = document.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_para.add_run(f"Report generated on: {timezone.now().strftime('%Y-%m-%d')}")
    
        range_para = document.add_paragraph()
        range_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        range_para.add_run(f"Violations from: {start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}")

        # Add spacing
        document.add_paragraph()

        # Summary Statistics
        document.add_paragraph(f"Total Violations: {violations.count()}")
    
        severity_para = document.add_paragraph()
        severity_para.add_run("Violations by Severity:").bold = True
    
        severity_counts = violations.values('violation__severity').annotate(count=models.Count('id'))
        for severity_count in severity_counts:
            severity = dict(Violation.SEVERITY_CHOICES).get(severity_count['violation__severity'], "Unknown")
            document.add_paragraph(f"{severity}: {severity_count['count']}", style='List Bullet')

        # Add spacing
        document.add_paragraph()

        # Detailed Records
        for record in violations:
            # Create a section for each violation record
            student_name = f"{record.student.first_name} {record.student.last_name}" if record.student else "Unknown"
            violation_name = record.violation.name if record.violation else "No Violation"
            date_recorded = record.recorded_at.strftime('%Y-%m-%d')
            severity = record.violation.get_severity_display() if record.violation else "Unknown"
            recorded_by = record.recorded_by.get_full_name() if record.recorded_by else "Unknown"
            recorded_at = record.recorded_at.strftime('%Y-%m-%d %H:%M:%S')

            record_section = document.add_paragraph()
            record_section.add_run("Student: ").bold = True
            record_section.add_run(student_name)
        
            record_section = document.add_paragraph()
            record_section.add_run("Violation: ").bold = True
            record_section.add_run(violation_name)
        
            record_section = document.add_paragraph()
            record_section.add_run("Date: ").bold = True
            record_section.add_run(date_recorded)
        
            record_section = document.add_paragraph()
            record_section.add_run("Severity: ").bold = True
            record_section.add_run(severity)
        
            record_section = document.add_paragraph()
            record_section.add_run("Recorded By: ").bold = True
            record_section.add_run(recorded_by)
        
            record_section = document.add_paragraph()
            record_section.add_run("Recorded At: ").bold = True
            record_section.add_run(recorded_at)

            # Add spacing between records
            document.add_paragraph()

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        document.save(response)
        messages.success(request, f"{title} has been exported as a Word file.")
        return response
        
    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "sanction":
            # If a violation is selected, filter sanctions by that violation
            if 'violation' in request.GET:
                kwargs["queryset"] = Sanction.objects.filter(violation_id=request.GET['violation'])
        return super().formfield_for_foreignkey(db_field, request, **kwargs)

    def add_user_feedback(self, request, obj, message, level=messages.INFO):
        """Utility method to add user feedback in the admin interface."""
        if not hasattr(request, '_messages'):  # Ensure messages framework is available
            return
        messages.add_message(request, level, message)

    def save_model(self, request, obj, form, change):
        super().save_model(request, obj, form, change)
        feedback_message = f"{'Updated' if change else 'Added'} record for {obj.student.first_name} {obj.student.last_name}."
        self.add_user_feedback(request, obj, feedback_message)

    def delete_model(self, request, obj):
        feedback_message = f"Deleted record for {obj.student.first_name} {obj.student.last_name}."
        self.add_user_feedback(request, obj, feedback_message, level=messages.WARNING)
        super().delete_model(request, obj)

    export_form_class = ExportForm

    def view_qr_code(self, obj):
        try:
            # Access the QR code from the related StudentRegistration model
            student_registration = obj.student.studentregistration  # Adjust this based on actual field name
            if student_registration and student_registration.qr_code:
                return student_registration.qr_code
        except (AttributeError, ObjectDoesNotExist):
            pass
        return "No QR Code"
    view_qr_code.short_description = 'QR Code'

    def get_queryset(self, request):
        # Use select_related to optimize query by including the student and studentregistration
        return super().get_queryset(request).select_related(
            'student',
            'student__studentregistration'  # Make sure this matches the related_name of the relationship
        )

    def status(self, obj):
        # Return a human-readable status based on the record's data
        if obj.sanction:
            remaining_hours = max(0, obj.sanction.duration_value - obj.total_hours_complied)
            return "Completed" if remaining_hours <= 0 else "In Progress"
        return "No Sanction Assigned"

    status.short_description = 'Status'  # Set the column header for the admin panel

    #def actions_buttons(self, obj):
        #assign_qr_url = reverse('admin:adminmod_violationrecord_assign_qr_code', args=[obj.pk])
        #time_in_url = reverse('admin:adminmod_violationrecord_time_in', args=[obj.pk])
        #time_out_url = reverse('admin:adminmod_violationrecord_time_out', args=[obj.pk])
        #return format_html(
            #'<a class="button" href="{}">Assign QR Code</a> ',
            #'<a class="button" href="{}">Time In</a> '
            #'<a class="button" href="{}">Time Out</a>',
           #assign_qr_url
       # )

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'assign_qr_code/<int:record_id>/',
                self.admin_site.admin_view(self.assign_qr_code),
                name='adminmod_violationrecord_assign_qr_code'  # Updated name following Django conventions
            ),
            path(
                'time_in/<int:record_id>/',
                self.admin_site.admin_view(self.time_in),
                name='adminmod_violationrecord_time_in'  # Updated name following Django conventions
            ),
            path(
                'time_out/<int:record_id>/',
                self.admin_site.admin_view(self.time_out),
                name='adminmod_violationrecord_time_out'  # Updated name following Django conventions
            ),
        ]
        return custom_urls + urls

    def assign_qr_code(self, request, record_id):
        record = get_object_or_404(ViolationRecord, pk=record_id)
        if request.method == "POST":
            qr_code = request.POST.get('qr_code')
            if qr_code:
                record.qr_code = qr_code
                record.save()
                self.message_user(request, "QR Code assigned successfully!")
                return HttpResponse(
                    '<script>window.opener.location.reload(); window.close();</script>'
                )
            else:
                self.message_user(request, "QR Code not provided", level='ERROR')
                return HttpResponse(
                    f'''
                    <form method="POST">
                        <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                        <p style="color: red;">QR Code is required. Please enter a QR Code.</p>
                        <input type="text" name="qr_code" required>
                        <input type="submit" value="Assign">
                    </form>
                    '''
                )
        
        return HttpResponse(
            f'''
            <form method="POST">
                <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                <p>Please enter QR code to assign:</p>
                <input type="text" name="qr_code" required autofocus>
                <input type="submit" value="Assign">
            </form>
            '''
        )

    def assign_qr_code_button(self, obj):
        return format_html(
            '<a class="button" onclick="window.open(\'assign_qr_code/{}/\', \'Assign QR Code\', \'width=400,height=200\')" href="javascript:void(0)">Assign QR Code</a>',
            obj.id
        )
    assign_qr_code_button.short_description = 'Assign QR Code'
    assign_qr_code_button.allow_tags = True

    

    def time_in(self, request, record_id):
        record = get_object_or_404(ViolationRecord, pk=record_id)
        if request.method == "POST":
            qr_code = request.POST.get('qr_code')
            if qr_code == record.qr_code:
                record.time_in = now()
                record.save()
                self.message_user(request, "Time In recorded successfully!")
                return HttpResponse(
                    '<script>window.opener.location.reload(); window.close();</script>'
                )
            else:
                self.message_user(request, "Invalid QR Code", level='ERROR')
                return HttpResponse(
                    f'''
                    <form method="POST">
                        <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                        <p style="color: red;">Invalid QR Code. Please try again.</p>
                        <input type="text" name="qr_code" required>
                        <input type="submit" value="Submit">
                    </form>
                    '''
                )
        
        return HttpResponse(
            f'''
            <form method="POST">
                <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                <p>Please scan or enter QR code:</p>
                <input type="text" name="qr_code" required autofocus>
                <input type="submit" value="Submit">
            </form>
            '''
        )


    def time_in_button(self, obj):
        return format_html(
            '<a onclick="window.open(\'time_in/{}/\', \'Time In\', \'width=400,height=200\')" '
            'href="javascript:void(0)" style="color: #4CAF50; text-decoration: underline; font-size: 14px;">Time In</a>',
            obj.id
        )
    time_in_button.short_description = 'Time In'
    time_in_button.allow_tags = True



    def time_out(self, request, record_id):
        record = get_object_or_404(ViolationRecord, pk=record_id)
        if request.method == "POST":
            qr_code = request.POST.get('qr_code')
            if qr_code == record.qr_code:
                record.time_out = now()
                record.update_total_hours()
                self.message_user(request, "Time Out recorded and hours updated!")
                return HttpResponse(
                    '<script>window.opener.location.reload(); window.close();</script>'
                )
            else:
                self.message_user(request, "Invalid QR Code", level='ERROR')
                return HttpResponse(
                    f'''
                    <form method="POST">
                        <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                        <p style="color: red;">Invalid QR Code. Please try again.</p>
                        <input type="text" name="qr_code" required>
                        <input type="submit" value="Submit">
                    </form>
                    '''
                )
        
        return HttpResponse(
            f'''
            <form method="POST">
                <input type="hidden" name="csrfmiddlewaretoken" value="{csrf.get_token(request)}">
                <p>Please scan or enter QR code:</p>
                <input type="text" name="qr_code" required autofocus>
                <input type="submit" value="Submit">
            </form>
            '''
        )

    def time_out_button(self, obj):
        return format_html(
            '<a class="button" onclick="window.open(\'time_out/{}/\', \'Time Out\', \'width=400,height=200\')" href="javascript:void(0)"style="color: #FF0000; text-decoration: underline; font-size: 14px;">Time Out</a>',
            obj.id
        )
    time_out_button.short_description = 'Time Out'
    time_out_button.allow_tags = True